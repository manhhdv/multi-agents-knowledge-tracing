#!/usr/bin/env python3
"""Blind human rating of the Feedback Generator (30 MathDial cases, two raters).

    make     bundle directory       -> rater_A.xlsx, rater_B.xlsx, HUONG_DAN_CHAM_FEEDBACK.md
    analyze  rater_A.xlsx B.xlsx    -> feedback_human_report.txt, feedback_human_metrics.csv

Raters see exactly the content the LLM judge saw (dialogue, student turn, reference answer, the five feedback
fields). They never see the judge's scores, the SAVA verdict, the mastery band or whether the output fell back
to the rule-based template. Case order is shuffled with a fixed seed so verdict groups cannot be inferred.
The human mean is the primary score; the LLM judge is reported for reference with its agreement to humans.
"""
import argparse
import glob
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from scipy.stats import spearmanr, wilcoxon
from sklearn.metrics import cohen_kappa_score

CRITERIA = ["correctness", "relevance", "clarity"]
FEEDBACK_KEYS = ["feedback_text", "scaffolding_question", "mastery_adaptation", "pedagogical_strategy", "next_step_hint"]
RATE_SHEET = "Cham diem"
GUIDE_SHEET = "Huong dan"
SEED = 221

GUIDELINE = """HƯỚNG DẪN CHẤM FEEDBACK CỦA HỆ THỐNG GIA SƯ — 30 trường hợp (MathDial)

MỤC TIÊU
Mỗi dòng gồm một đoạn hội thoại gia sư – học sinh, lượt trả lời mới nhất của học sinh, đáp án đúng của bài,
và phản hồi (feedback) do hệ thống sinh ra cho lượt đó. Chấm chất lượng feedback theo 3 tiêu chí, thang 1–5.

QUY TẮC BẮT BUỘC
1. Chấm độc lập: không trao đổi với người chấm còn lại cho đến khi cả hai nộp file.
2. Không dùng ChatGPT/Claude/công cụ AI nào để hỗ trợ chấm.
3. Chỉ điền các cột chấm điểm và ghi chú. Không sửa, không xóa, không sắp xếp lại dòng.
4. Tự đánh giá câu trả lời của học sinh đúng hay sai dựa trên hội thoại và đáp án; hệ thống không cho bạn biết.
5. Chấm từng tiêu chí riêng: một feedback rõ ràng nhưng sai toán vẫn có thể được Clarity cao, Correctness thấp.

3 TIÊU CHÍ (giống định nghĩa đưa cho LLM chấm)
• correctness — Toán học trong feedback đúng, và nhận định về câu trả lời của học sinh (đúng/sai/chưa rõ) là đúng.
• relevance — Feedback nói đúng vào lượt này của học sinh, lỗi của học sinh và kỹ năng yếu nhất, không chung chung.
• clarity — Feedback rõ ràng, dễ hiểu và phù hợp với người học.

THANG ĐIỂM
  5 — Rất tốt: không có vấn đề nào.
  4 — Tốt: một vấn đề nhỏ, không ảnh hưởng đến việc học.
  3 — Tạm được: có vấn đề rõ ràng nhưng feedback vẫn dùng được.
  2 — Kém: vấn đề nghiêm trọng, dễ gây hiểu nhầm hoặc vô ích.
  1 — Rất kém: sai hoàn toàn, lạc đề hoặc gây hại cho việc học.

VÍ DỤ ĐỊNH MỨC CORRECTNESS (không lấy từ dữ liệu)
  Học sinh: "3/4 + 1/4 = 4/8". Đáp án: 1.
  5 — "Gần đúng rồi! Khi cộng hai phân số cùng mẫu, mẫu số có thay đổi không?"
  3 — "Chưa đúng. Hãy kiểm tra lại cách cộng phân số." (đúng nhưng không chỉ ra lỗi)
  1 — "Chính xác, 4/8 là kết quả đúng!" (nhận định sai)

CỘT lo_dap_an (có / khong)
Chọn "co" nếu feedback nói ra đáp án cuối cùng khi học sinh CHƯA trả lời đúng. Nếu học sinh đã đúng, chọn "khong".

CỘT ghi_chu (không bắt buộc)
Ghi ngắn lý do khi cho điểm 1–2 hoặc khi phân vân.

KHI XONG
Kiểm tra ô "Đã chấm" bên dưới báo 30/30, lưu file giữ nguyên tên, gửi lại cho người điều phối.
"""


def load_cases(bundle):
    bundle = Path(bundle)
    cases = pd.read_csv(bundle / "feedback_generator_cases.csv")
    qual = pd.read_csv(glob.glob(str(bundle / "mathdial/qlora/seed221/test/qual_*.csv"))[0])
    qual = qual[qual["Dialogue ID"].notna()].copy()
    qual["Dialogue ID"] = qual["Dialogue ID"].astype(int)
    qual["Turn"] = qual["Turn"].astype(int)
    raw = {(int(r["dialogue"]), int(r["turn_id"])): r
           for r in json.loads((bundle / "mathdial/qlora/seed221/test/raw_logit_gaps.json").read_text())}
    records = []
    for case in cases.itertuples(index=False):
        dialogue = qual[qual["Dialogue ID"] == case.dialogue].sort_values("Turn")
        lines = []
        for turn in dialogue[dialogue.Turn < case.turn_id].itertuples(index=False):
            if str(turn.Teacher) not in ("--", "nan"):
                lines.append(f"Tutor: {turn.Teacher}")
            lines.append(f"Student: {turn.Student}")
        current = dialogue[dialogue.Turn == case.turn_id]
        assert len(current) == 1, (case.dialogue, case.turn_id)
        lines.append(f"Tutor: {current.Teacher.iloc[0]}")
        row = raw[(int(case.dialogue), int(case.turn_id))]
        assert str(current.Student.iloc[0]).strip() == str(row["student_text"]).strip(), (case.dialogue, case.turn_id)
        feedback = "\n\n".join(f"[{key}]\n{getattr(case, key)}" for key in FEEDBACK_KEYS)
        records.append({"dialogue": int(case.dialogue), "turn_id": int(case.turn_id), "context": "\n".join(lines),
                        "student_text": row["student_text"], "reference": row["reference"], "feedback": feedback})
    frame = pd.DataFrame(records)
    return frame.sample(frac=1, random_state=SEED).reset_index(drop=True)


def dropdown(ws, column, options, n):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False,
                        showErrorMessage=True, errorTitle="Giá trị không hợp lệ", error="Chọn: " + ", ".join(options))
    ws.add_data_validation(dv)
    dv.add(f"{column}2:{column}{n + 1}")


def cmd_make(args):
    frame = load_cases(args.bundle)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "HUONG_DAN_CHAM_FEEDBACK.md").write_text(GUIDELINE, encoding="utf-8")
    n = len(frame)
    for rater in args.raters:
        wb = Workbook()
        guide = wb.active
        guide.title = GUIDE_SHEET
        lines = GUIDELINE.splitlines()
        for i, line in enumerate(lines, start=1):
            guide.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
        guide.column_dimensions["A"].width = 125
        guide["A1"].font = Font(bold=True, size=13)
        guide.cell(row=len(lines) + 2, column=1, value=f"Người chấm: {rater}").font = Font(bold=True)
        guide.cell(row=len(lines) + 3, column=1,
                   value=f'="Đã chấm: "&COUNTIFS(\'{RATE_SHEET}\'!H2:H{n + 1},"<>",\'{RATE_SHEET}\'!I2:I{n + 1},"<>",'
                         f'\'{RATE_SHEET}\'!J2:J{n + 1},"<>",\'{RATE_SHEET}\'!K2:K{n + 1},"<>")&"/{n}"').font = Font(bold=True)

        ws = wb.create_sheet(RATE_SHEET)
        ws.append(["stt", "dialogue", "turn_id", "hoi_thoai (trước lượt này)", "luot_hoc_sinh", "dap_an",
                   "feedback_he_thong", "correctness", "relevance", "clarity", "lo_dap_an", "ghi_chu"])
        for i, row in enumerate(frame.itertuples(index=False), start=1):
            ws.append([i, row.dialogue, row.turn_id, row.context, row.student_text, row.reference, row.feedback,
                       None, None, None, None, None])
        widths = {"A": 5, "B": 9, "C": 8, "D": 60, "E": 35, "F": 9, "G": 60, "H": 12, "I": 11, "J": 10, "K": 11, "L": 28}
        rating_fill = PatternFill("solid", fgColor="E8F5E9")
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDDDDD")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.column_letter in "HIJKL":
                    cell.fill = rating_fill
        for column in "HIJ":
            dropdown(ws, column, ["1", "2", "3", "4", "5"], n)
        dropdown(ws, "K", ["co", "khong"], n)
        ws.freeze_panes = "A2"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
        wb.active = 1
        path = out / f"rater_{rater}.xlsx"
        wb.save(path)
        print("wrote", path)


def read_ratings(path):
    frame = pd.read_excel(path, sheet_name=RATE_SHEET)
    problems = []
    for criterion in CRITERIA:
        values = pd.to_numeric(frame[criterion], errors="coerce")
        problems += frame.loc[~values.isin([1, 2, 3, 4, 5]), "stt"].tolist()
        frame[criterion] = values
    frame["lo_dap_an"] = frame.lo_dap_an.fillna("").astype(str).str.strip().str.lower()
    problems += frame.loc[~frame.lo_dap_an.isin(["co", "khong"]), "stt"].tolist()
    if problems:
        sys.exit(f"{path}: rows with a missing or invalid rating: {sorted(set(problems))}")
    return frame


def cmd_analyze(args):
    a, b = read_ratings(args.a), read_ratings(args.b)
    keys = ["dialogue", "turn_id"]
    merged = a.merge(b[keys + CRITERIA + ["lo_dap_an"]], on=keys, suffixes=("_A", "_B"))
    cases = pd.read_csv(Path(args.bundle) / "feedback_generator_cases.csv")
    merged = merged.merge(cases[keys + ["verdict", "source", "answer_leak"] + CRITERIA]
                          .rename(columns={c: f"{c}_judge" for c in CRITERIA}), on=keys)
    assert len(merged) == len(a) == len(b) == len(cases), "Rater files and cases do not cover the same turns."

    rows, lines = [], [f"n = {len(merged)} cases, 2 raters (primary score = mean of the two raters)\n"]
    for criterion in CRITERIA:
        ra, rb, judge = merged[f"{criterion}_A"], merged[f"{criterion}_B"], merged[f"{criterion}_judge"]
        human = (ra + rb) / 2
        kappa_w = cohen_kappa_score(ra, rb, weights="quadratic", labels=[1, 2, 3, 4, 5])
        rho, rho_p = spearmanr(human, judge)
        diff = judge - human
        try:
            bias_p = wilcoxon(judge, human).pvalue
        except ValueError:
            bias_p = float("nan")
        row = {"criterion": criterion, "human_mean": human.mean(), "human_sd": human.std(ddof=1),
               "rater_A_mean": ra.mean(), "rater_B_mean": rb.mean(),
               "weighted_kappa_A_B": kappa_w, "exact_agreement_A_B": float((ra == rb).mean()),
               "within_1_agreement_A_B": float(((ra - rb).abs() <= 1).mean()),
               "judge_mean_reference": judge.mean(), "spearman_judge_vs_human": rho, "spearman_p": rho_p,
               "judge_minus_human_mean": diff.mean(), "wilcoxon_p_judge_vs_human": bias_p}
        for verdict in ["correct", "incorrect", "undetermined"]:
            row[f"human_mean_{verdict}"] = human[merged.verdict == verdict].mean()
            row[f"judge_mean_{verdict}"] = judge[merged.verdict == verdict].mean()
        rows.append(row)
    metrics = pd.DataFrame(rows)
    metrics.to_csv(Path(args.out) / "feedback_human_metrics.csv", index=False)
    lines.append(metrics.round(3).T.to_string())

    leak_a, leak_b = merged.lo_dap_an_A == "co", merged.lo_dap_an_B == "co"
    not_correct = merged.verdict != "correct"
    auto = merged.answer_leak[not_correct].astype(float)
    human_leak = (leak_a | leak_b)[not_correct]
    lines.append(f"\nAnswer leakage on turns not verified correct (n={int(not_correct.sum())}):"
                 f"\n  rater A {leak_a[not_correct].mean():.3f}, rater B {leak_b[not_correct].mean():.3f}, "
                 f"either {human_leak.mean():.3f}; Cohen's kappa A-B "
                 f"{cohen_kappa_score(leak_a[not_correct], leak_b[not_correct]):.3f}"
                 f"\n  automatic string-match check {auto.mean():.3f}; agreement with 'either rater' "
                 f"{float(((auto == 1) == human_leak).mean()):.3f}")
    lines.append(f"\nHuman mean by output source (llm vs rule_fallback):\n" + merged.assign(
        **{f"{c}_human": (merged[f"{c}_A"] + merged[f"{c}_B"]) / 2 for c in CRITERIA}).groupby("source")[
        [f"{c}_human" for c in CRITERIA]].mean().round(2).to_string())
    report = "\n".join(lines)
    (Path(args.out) / "feedback_human_report.txt").write_text(report, encoding="utf-8")
    print(report)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    make = sub.add_parser("make")
    make.add_argument("--bundle", required=True)
    make.add_argument("--out", default=".")
    make.add_argument("--raters", nargs="+", default=["A", "B"])
    analyze = sub.add_parser("analyze")
    analyze.add_argument("a")
    analyze.add_argument("b")
    analyze.add_argument("--bundle", required=True)
    analyze.add_argument("--out", default=".")
    args = parser.parse_args()
    {"make": cmd_make, "analyze": cmd_analyze}[args.command](args)


if __name__ == "__main__":
    main()
